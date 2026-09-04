"""Extract raw text from uploaded documents."""
import io

from docx import Document as DocxDocument
from pypdf import PdfReader


class ParsingError(Exception):
    pass


def extract_text(file_bytes: bytes, content_type: str, filename: str) -> str:
    lower_name = filename.lower()

    if content_type == "application/pdf" or lower_name.endswith(".pdf"):
        return _extract_pdf(file_bytes)

    if (
        content_type
        in (
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        or lower_name.endswith(".docx")
    ):
        return _extract_docx(file_bytes)

    raise ParsingError(f"Unsupported file type: {content_type} ({filename})")


def _extract_pdf(file_bytes: bytes) -> str:
    try:
        reader = PdfReader(io.BytesIO(file_bytes))
        pages = []
        for page in reader.pages:
            pages.append(page.extract_text() or "")
        text = "\n\n".join(pages).strip()
    except Exception as exc:  # noqa: BLE001
        raise ParsingError(f"Failed to parse PDF: {exc}") from exc

    if not text:
        raise ParsingError("No extractable text found in PDF (it may be scanned/image-only).")
    return text


def _extract_docx(file_bytes: bytes) -> str:
    try:
        doc = DocxDocument(io.BytesIO(file_bytes))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        parts.append(cell.text)
        text = "\n".join(parts).strip()
    except Exception as exc:  # noqa: BLE001
        raise ParsingError(f"Failed to parse DOCX: {exc}") from exc

    if not text:
        raise ParsingError("No extractable text found in DOCX.")
    return text
